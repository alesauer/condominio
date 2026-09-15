from datetime import date
from io import BytesIO
from fastapi import APIRouter, Depends, Query, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from app.core.database import get_db
from app.core.permissions import admin_required
from app.api.deps import get_current_user
from app.models.receita import Receita
from app.models.despesa import Despesa

router = APIRouter()


@router.get("/balancete", dependencies=[Depends(get_current_user)])
async def balancete(
    mes: int = Query(..., ge=1, le=12),
    ano: int = Query(..., ge=2000, le=2100),
    formato: str = Query("excel"),
    db: AsyncSession = Depends(get_db),
):
    inicio = date(ano, mes, 1)
    if mes == 12:
        fim = date(ano + 1, 1, 1)
    else:
        fim = date(ano, mes + 1, 1)

    receitas_r = await db.execute(
        select(Receita)
        .where(Receita.competencia >= inicio, Receita.competencia < fim)
        .order_by(Receita.created_at)
    )
    despesas_r = await db.execute(
        select(Despesa)
        .where(Despesa.competencia >= inicio, Despesa.competencia < fim)
        .order_by(Despesa.created_at)
    )

    receitas = receitas_r.scalars().all()
    despesas = despesas_r.scalars().all()

    total_receitas = sum(float(r.valor) for r in receitas)
    total_despesas = sum(float(d.valor) for d in despesas)
    saldo = total_receitas - total_despesas

    if formato == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = f"Balancete {mes:02d}-{ano}"

        # Cabeçalho Principal
        ws.merge_cells("A1:D1")
        ws["A1"] = f"CONDO GESTÃO — BALANCETE MENSAL ({mes:02d}/{ano})"
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Cabeçalhos da Tabela
        headers = ["Tipo", "Descrição / Categoria", "Status", "Valor (R$)"]
        ws.append(headers)
        ws.row_dimensions[2].height = 20
        for col_num in range(1, 5):
            cell = ws.cell(row=2, column=col_num)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
            cell.alignment = Alignment(horizontal="center" if col_num != 2 else "left")

        # Linhas de Receitas
        for r in receitas:
            ws.append(["Receita", r.descricao, str(r.status.value if hasattr(r.status, 'value') else r.status), float(r.valor)])

        # Linhas de Despesas
        for d in despesas:
            ws.append(["Despesa", d.descricao, str(d.status.value if hasattr(d.status, 'value') else d.status), -float(d.valor)])

        # Linha em branco
        ws.append([])

        # Resumo
        ws.append(["", "TOTAL RECEITAS", "", total_receitas])
        ws.append(["", "TOTAL DESPESAS", "", total_despesas])
        ws.append(["", "SALDO LÍQUIDO", "", saldo])

        # Ajuste de larguras
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 20

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=balancete_{mes:02d}_{ano}.xlsx"},
        )

    elif formato == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#0f172a"),
            alignment=1, # Centralizado
        )
        subtitle_style = ParagraphStyle(
            name="SubtitleStyle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748b"),
            alignment=1,
        )
        cell_style = ParagraphStyle(
            name="CellStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
        )
        header_cell_style = ParagraphStyle(
            name="HeaderCellStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            fontName="Helvetica-Bold",
            textColor=colors.white,
        )

        elements = []
        elements.append(Paragraph("<b>Condo Gestão — Sistema Residencial</b>", title_style))
        elements.append(Paragraph(f"Demonstrativo de Balancete Mensal • Competência: {mes:02d}/{ano}", subtitle_style))
        elements.append(Spacer(1, 16))

        # Dados da Tabela
        table_data = [
            [
                Paragraph("<b>Tipo</b>", header_cell_style),
                Paragraph("<b>Descrição</b>", header_cell_style),
                Paragraph("<b>Status</b>", header_cell_style),
                Paragraph("<b>Valor (R$)</b>", header_cell_style),
            ]
        ]

        for r in receitas:
            st = str(r.status.value if hasattr(r.status, 'value') else r.status)
            table_data.append([
                Paragraph("Receita", cell_style),
                Paragraph(r.descricao, cell_style),
                Paragraph(st, cell_style),
                Paragraph(f"R$ {float(r.valor):.2f}", cell_style),
            ])

        for d in despesas:
            st = str(d.status.value if hasattr(d.status, 'value') else d.status)
            table_data.append([
                Paragraph("Despesa", cell_style),
                Paragraph(d.descricao, cell_style),
                Paragraph(st, cell_style),
                Paragraph(f"- R$ {float(d.valor):.2f}", cell_style),
            ])

        if len(table_data) == 1:
            table_data.append([
                Paragraph("—", cell_style),
                Paragraph("Nenhum lançamento registrado nesta competência.", cell_style),
                Paragraph("—", cell_style),
                Paragraph("R$ 0,00", cell_style),
            ])

        t = Table(table_data, colWidths=[70, 260, 90, 100])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("ALIGN", (3, 0), (3, -1), "RIGHT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#ffffff"), colors.HexColor("#f8fafc")]),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 16))

        # Tabela de Resumo Financeiro
        summary_data = [
            [Paragraph("<b>Total de Receitas:</b>", cell_style), Paragraph(f"<b>R$ {total_receitas:.2f}</b>", cell_style)],
            [Paragraph("<b>Total de Despesas:</b>", cell_style), Paragraph(f"<b>R$ {total_despesas:.2f}</b>", cell_style)],
            [Paragraph("<b>Saldo Operacional:</b>", cell_style), Paragraph(f"<b>R$ {saldo:.2f}</b>", cell_style)],
        ]
        sum_t = Table(summary_data, colWidths=[200, 150], hAlign="RIGHT")
        sum_t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f5f9")),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(sum_t)

        doc.build(elements)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=balancete_{mes:02d}_{ano}.pdf"},
        )

    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Formato inválido. Use 'excel' ou 'pdf'.")
